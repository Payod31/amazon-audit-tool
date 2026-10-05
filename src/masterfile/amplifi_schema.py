"""
Reads a Pattern "Amplifi" import file -- the company's own product-data
import template, distinct from Amazon's masterfile but built from the same
underlying Amazon attribute vocabulary (its column headers are literally
Amazon attribute paths, just with a language suffix like " - en-US").

An Amplifi sheet is always a SUBSET of a product type's full Amazon
masterfile attributes, plus a couple of Amplifi-only operational columns
(Collection Folder, Import Type, ...) that are not product data at all --
those are filled by the user's own team, never by us.

This cross-references each Amplifi attribute column against the Amazon
masterfile schema already parsed and stored (see schema.py) for the
product type this file is for, so labels / required level / dropdown
values don't need to be rediscovered -- only genuinely new attributes
(ones no stored masterfile knows about) get flagged for the user.
"""

from __future__ import annotations

import io
import json
import os
import re
import time
from dataclasses import asdict, dataclass, field

import openpyxl

from src.masterfile.schema import load_schema

DATA_SHEET_CANDIDATES = ["Draft and Approved"]
DEFINITIONS_SHEET_CANDIDATES = ["Data DefinitionsValid Values", "Data Definitions"]
IMPORT_TYPE_SHEET = "Import Type"

LANGUAGE_SUFFIX_RE = re.compile(r"\s*-\s*[a-z]{2}-[A-Z]{2}\s*$")

STORAGE_DIR = os.path.join("storage", "amplifi_schemas")


def _strip_language_suffix(header: str) -> str:
    return LANGUAGE_SUFFIX_RE.sub("", header or "").strip()


def _detect_header_row(ws, max_scan: int = 6) -> int:
    """The real field-label row is the one with the most non-blank cells
    among the first few rows -- a group-marker row above it (like
    'Required' / 'Attributes (add Attribute Lables below)') usually has
    only a handful of values, not one per column, so picking by count
    (rather than 'first row with >= N values') tells them apart."""
    best_row, best_count = 1, -1
    for r in range(1, min(max_scan, ws.max_row) + 1):
        count = sum(1 for c in range(1, ws.max_column + 1) if ws.cell(row=r, column=c).value)
        if count > best_count:
            best_row, best_count = r, count
    if best_count <= 0:
        raise ValueError("Could not find a header row with any labels in the first few rows.")
    return best_row


def _find_data_sheet(wb):
    for name in DATA_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return wb[name]
    # Fall back: the sheet with the most columns is almost certainly the
    # data-entry sheet in this kind of file.
    return max(wb.worksheets, key=lambda ws: ws.max_column)


@dataclass
class SystemFieldDef:
    """An Amplifi-only operational column (Collection Folder, Import Type,
    ...) -- not product data, always filled by the user's own team."""
    label: str
    description: str = ""
    accepted_values: str = ""
    example: str = ""
    required: str = ""
    valid_values: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "SystemFieldDef":
        return cls(**d)


@dataclass
class AttributeColumn:
    column_index: int
    header: str  # raw header exactly as it appears in the file
    attribute_path: str  # header with the language suffix stripped
    matched: bool = False
    label: str = ""
    group: str = ""
    required_level: str = ""
    dropdown_values: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "AttributeColumn":
        return cls(**d)


@dataclass
class AmplifiSchema:
    source_filename: str
    sheet_name: str
    header_row: int
    data_row: int
    product_types: list
    system_fields: list  # list[SystemFieldDef]
    attribute_columns: list  # list[AttributeColumn]
    unmatched_attributes: list  # attribute paths not found in any stored masterfile
    parsed_at: float = field(default_factory=time.time)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "AmplifiSchema":
        d = dict(d)
        d["system_fields"] = [SystemFieldDef.from_dict(x) for x in d["system_fields"]]
        d["attribute_columns"] = [AttributeColumn.from_dict(x) for x in d["attribute_columns"]]
        return cls(**d)


def parse_amplifi_template(data: bytes, filename: str) -> AmplifiSchema:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = _find_data_sheet(wb)

    header_row = _detect_header_row(ws)
    data_row = header_row + 1

    # Which columns are plain Amplifi system fields (no attribute-path
    # punctuation, no language suffix) vs. real attribute columns?
    system_fields_cols = []
    attribute_cols = []
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=header_row, column=c).value
        if not header:
            continue
        header = str(header).strip()
        stripped = _strip_language_suffix(header)
        looks_like_attribute = stripped != header or any(ch in header for ch in "#[]=")
        if looks_like_attribute:
            attribute_cols.append((c, header, stripped))
        else:
            system_fields_cols.append((c, header))

    # Definitions for the system fields, from the Data Definitions sheet
    # (keyed by the plain label, e.g. "Collection Folder").
    definitions: dict[str, tuple] = {}
    for name in DEFINITIONS_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            dd = wb[name]
            for r in range(1, dd.max_row + 1):
                label = dd.cell(row=r, column=1).value
                if not label:
                    continue
                desc = dd.cell(row=r, column=2).value or ""
                accepted = dd.cell(row=r, column=3).value or ""
                example = dd.cell(row=r, column=4).value or ""
                required = dd.cell(row=r, column=5).value or ""
                definitions[str(label).strip()] = (str(desc).strip(), str(accepted).strip(), str(example).strip(), str(required).strip())
            break

    import_type_values = []
    if IMPORT_TYPE_SHEET in wb.sheetnames:
        it = wb[IMPORT_TYPE_SHEET]
        for r in range(2, it.max_row + 1):
            v = it.cell(row=r, column=1).value
            if v:
                import_type_values.append(str(v).strip())

    system_fields = []
    for c, label in system_fields_cols:
        desc, accepted, example, required = definitions.get(label, ("", "", "", ""))
        valid_values = import_type_values if label.lower() == "import type" else []
        system_fields.append(SystemFieldDef(
            label=label, description=desc, accepted_values=accepted,
            example=example, required=required, valid_values=valid_values,
        ))

    # Figure out the product type(s) this file's data rows are for, by
    # reading whichever attribute column is "product_type#1.value".
    product_types: list[str] = []
    pt_col = next((c for c, h, s in attribute_cols if s.startswith("product_type")), None)
    if pt_col:
        seen = set()
        for r in range(data_row, ws.max_row + 1):
            v = ws.cell(row=r, column=pt_col).value
            if v:
                v = str(v).strip()
                if v and v not in seen:
                    seen.add(v)
                    product_types.append(v)

    # Cross-reference every attribute column against the stored masterfile
    # schema(s) for the product type(s) this file is for.
    stored_schemas = {pt: load_schema(pt) for pt in product_types}
    attribute_columns: list[AttributeColumn] = []
    unmatched: list[str] = []
    for c, header, attr_path in attribute_cols:
        col = AttributeColumn(column_index=c, header=header, attribute_path=attr_path)
        found = False
        for pt, schema in stored_schemas.items():
            if schema is None:
                continue
            match = next((f for f in schema.fields if f.attribute_path == attr_path), None)
            if match:
                col.matched = True
                col.label = match.label
                col.group = match.group
                col.required_level = match.required_level
                col.dropdown_values = match.values_for(pt)
                found = True
                break
        if not found:
            unmatched.append(attr_path)
        attribute_columns.append(col)

    return AmplifiSchema(
        source_filename=filename,
        sheet_name=ws.title,
        header_row=header_row,
        data_row=data_row,
        product_types=product_types,
        system_fields=system_fields,
        attribute_columns=attribute_columns,
        unmatched_attributes=unmatched,
    )


def _safe_filename(key: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]", "_", key) + ".json"


def _schema_path(key: str) -> str:
    return os.path.join(STORAGE_DIR, _safe_filename(key))


def load_amplifi_schema(key: str) -> AmplifiSchema | None:
    path = _schema_path(key)
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as fh:
        return AmplifiSchema.from_dict(json.load(fh))


def find_attributes_start_column(ws, header_row: int) -> int:
    """Finds the column where the 'Attributes (add Attribute Lables
    below)' marker sits in the row above the header row, for a blank
    Amplifi template that has no attribute columns filled in yet."""
    marker_row = max(1, header_row - 1)
    for c in range(1, ws.max_column + 50):
        v = ws.cell(row=marker_row, column=c).value
        # Look for the specific "add attribute(s) ... below" instruction,
        # not just any cell that happens to mention "attribute" (a group
        # header like "Variant Attributes" would otherwise match first).
        if v and "add attribute" in str(v).lower():
            return c
    raise ValueError(
        "Could not find the 'Attributes (add Attribute Lables below)' "
        "marker, so I don't know where the system columns end and the "
        "attribute columns should start."
    )


def fill_blank_template_with_known_row(
    template_bytes: bytes,
    attribute_values: list,  # list of (header_with_suffix, value) in display order
) -> bytes:
    """Takes a blank Amplifi template (system columns already labeled,
    Attributes section empty) and writes one data row into it: the
    attribute headers go into the header row starting at the 'Attributes'
    marker column, and the matching values go into the first data row
    right below. System columns (Product, Import Type, Roles, ...) are
    left exactly as they were -- untouched, still blank -- since those are
    filled by the user's own team, not from product data."""
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=False)
    ws = _find_data_sheet(wb)

    header_row = _detect_header_row(ws)
    data_row = header_row + 1

    start_col = find_attributes_start_column(ws, header_row)

    for i, (header, value) in enumerate(attribute_values):
        col = start_col + i
        ws.cell(row=header_row, column=col, value=header)
        ws.cell(row=data_row, column=col, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def save_amplifi_schema(schema: AmplifiSchema) -> dict:
    """Stores (or merges into) one Amplifi template shape per product
    type it covers. Returns a small report dict per product type."""
    os.makedirs(STORAGE_DIR, exist_ok=True)
    reports = {}
    keys = schema.product_types or ["UNKNOWN"]
    for key in keys:
        existing = load_amplifi_schema(key)
        if existing is None:
            new_schema = schema
            added = [c.header for c in schema.attribute_columns]
            is_new = True
        else:
            by_path = {c.attribute_path: c for c in existing.attribute_columns}
            merged_cols = []
            added = []
            for c in schema.attribute_columns:
                if c.attribute_path in by_path:
                    merged_cols.append(c)  # refresh with latest-known info
                    del by_path[c.attribute_path]
                else:
                    merged_cols.append(c)
                    added.append(c.header)
            merged_cols.extend(by_path.values())  # keep any not in this upload
            new_schema = AmplifiSchema(
                source_filename=schema.source_filename,
                sheet_name=schema.sheet_name,
                header_row=schema.header_row,
                data_row=schema.data_row,
                product_types=[key],
                system_fields=schema.system_fields or existing.system_fields,
                attribute_columns=merged_cols,
                unmatched_attributes=schema.unmatched_attributes,
            )
            is_new = False

        with open(_schema_path(key), "w", encoding="utf-8") as fh:
            json.dump(new_schema.to_dict(), fh, indent=2)

        reports[key] = {
            "is_new": is_new,
            "total_columns": len(new_schema.attribute_columns),
            "added": added,
        }
    return reports