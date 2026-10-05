"""
Takes a FILLED file (a real Amazon masterfile with product data already in
it, or a filled Amplifi import file) and copies its product-attribute data
into a BLANK template -- the same move done by hand for the Kreg router
bit example, generalized to any number of rows and either source shape.

The rule stays the same as every time this came up: only product-attribute
data gets copied. The blank template's own operational/system columns
(Collection Folder, Import Type, Roles, Region, ...) are never touched --
those are filled by the user's own team, not derived from product data.
"""

from __future__ import annotations

import io
import os

import openpyxl

from src.masterfile.amplifi_schema import (
    _detect_header_row,
    _find_data_sheet,
    _strip_language_suffix,
    find_attributes_start_column,
)
from src.masterfile.marketplaces import DEFAULT_MARKETPLACE_COUNTRY, locale_for_country
from src.masterfile.schema import _find_template_sheet, _parse_settings

# The blank target template is a fixed asset, not something re-uploaded
# every time -- the user hands it over once (it rarely changes) and from
# then on only a filled masterfile is needed.
DEFAULT_BLANK_TEMPLATE_PATH = os.path.join("templates", "data_import_template.xlsx")


def load_default_blank_template() -> bytes:
    """Reads the blank template stored alongside the app. Raises
    FileNotFoundError with a plain-English message if it isn't there yet."""
    if not os.path.exists(DEFAULT_BLANK_TEMPLATE_PATH):
        raise FileNotFoundError(
            f"No stored blank template found at '{DEFAULT_BLANK_TEMPLATE_PATH}'. "
            "Save the blank template there (or upload one to replace it) before filling."
        )
    with open(DEFAULT_BLANK_TEMPLATE_PATH, "rb") as fh:
        return fh.read()


def save_default_blank_template(data: bytes) -> None:
    """Replaces the stored blank template with a new one."""
    os.makedirs(os.path.dirname(DEFAULT_BLANK_TEMPLATE_PATH), exist_ok=True)
    with open(DEFAULT_BLANK_TEMPLATE_PATH, "wb") as fh:
        fh.write(data)


def detect_source_kind(wb) -> str:
    """'amazon_masterfile' if the workbook has Amazon's settings-string
    signature in a Template sheet's A1, otherwise 'amplifi'."""
    try:
        _find_template_sheet(wb)
        return "amazon_masterfile"
    except ValueError:
        return "amplifi"


def extract_filled_rows(data: bytes, filename: str) -> tuple[list[dict], str]:
    """Reads every filled data row out of a source file, as a list of
    {attribute_path: value} dicts (no language suffix, no system columns)
    -- one dict per product. Returns (rows, source_kind)."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    kind = detect_source_kind(wb)
    rows: list[dict] = []

    if kind == "amazon_masterfile":
        ws = _find_template_sheet(wb)
        settings = _parse_settings(ws.cell(row=1, column=1).value)
        attribute_row = int(settings.get("attributeRow", 5))
        data_row = int(settings.get("dataRow", 7))

        attr_by_col = {}
        for c in range(1, ws.max_column + 1):
            attr = ws.cell(row=attribute_row, column=c).value
            if attr:
                attr_by_col[c] = str(attr).strip()

        for r in range(data_row, ws.max_row + 1):
            row_dict = {}
            for c, attr in attr_by_col.items():
                v = ws.cell(row=r, column=c).value
                if v not in (None, ""):
                    row_dict[attr] = v
            if row_dict:
                rows.append(row_dict)

    else:  # amplifi
        ws = _find_data_sheet(wb)
        header_row = _detect_header_row(ws)
        data_row = header_row + 1

        attr_by_col = {}
        for c in range(1, ws.max_column + 1):
            header = ws.cell(row=header_row, column=c).value
            if not header:
                continue
            header = str(header).strip()
            stripped = _strip_language_suffix(header)
            looks_like_attribute = stripped != header or any(ch in header for ch in "#[]=")
            if looks_like_attribute:
                attr_by_col[c] = stripped

        for r in range(data_row, ws.max_row + 1):
            row_dict = {}
            for c, attr in attr_by_col.items():
                v = ws.cell(row=r, column=c).value
                if v not in (None, ""):
                    row_dict[attr] = v
            if row_dict:
                rows.append(row_dict)

    return rows, kind


def fill_blank_template(
    template_bytes: bytes,
    rows: list[dict],
    header_suffix: str | None = None,
) -> bytes:
    """Writes every row's attribute data into the blank template's
    Attributes section (one row per product, starting right after the
    header row), leaving every system/operational column exactly as it
    was -- untouched, still blank.

    header_suffix picks which Amazon marketplace/locale the attribute
    columns get labeled for (e.g. " - de-DE" for Germany, " - fr-FR" for
    France). Pass it explicitly -- the UI does, based on the country the
    user picks -- or leave it as None to fall back to the US locale."""
    if not rows:
        raise ValueError("No filled rows were found in the source file to copy over.")

    if header_suffix is None:
        header_suffix = f" - {locale_for_country(DEFAULT_MARKETPLACE_COUNTRY)}"

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=False)
    ws = _find_data_sheet(wb)

    header_row = _detect_header_row(ws)
    data_row = header_row + 1
    start_col = find_attributes_start_column(ws, header_row)

    # One shared column order: whatever order the first row's attributes
    # came in, then any further attributes seen only in later rows tacked
    # on at the end -- so every row lines up under the same headers.
    ordered_attrs: list[str] = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                ordered_attrs.append(key)

    for i, attr in enumerate(ordered_attrs):
        ws.cell(row=header_row, column=start_col + i, value=f"{attr}{header_suffix}")

    for r_offset, row in enumerate(rows):
        for i, attr in enumerate(ordered_attrs):
            value = row.get(attr)
            if value is not None:
                ws.cell(row=data_row + r_offset, column=start_col + i, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()