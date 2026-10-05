"""
Reads Amazon's "flat file" listing masterfile (the .xlsx/.xlsm you download
from Seller Central's bulk upload tool, fill in offline, and upload back
yourself) and extracts a reusable schema: every field's label, section,
required level, description/example, and -- critically -- the real
dropdown values Amazon wired into the file for that field and product
type.

Nothing here talks to Amazon or Seller Central. It only reads a file the
user already has on their machine and remembers what it found, so the
next masterfile (same product type, or a different one) can build on what
was learned instead of starting over.
"""

from __future__ import annotations

import io
import json
import os
import re
import time
import urllib.parse
from dataclasses import asdict, dataclass, field

import openpyxl
from openpyxl.utils import range_boundaries

TEMPLATE_SHEET_CANDIDATES = ["Template"]
DATA_DEFINITIONS_SHEET = "Data Definitions"

# Default row layout if the file's own settings string doesn't say otherwise
# (these match Amazon's current flat-file format as of this file).
DEFAULT_LABEL_ROW = 4
DEFAULT_ATTRIBUTE_ROW = 5
DEFAULT_DATA_ROW = 7

SANITIZE_RE = re.compile(r"[\[\]=#:\s]")

STORAGE_DIR = os.path.join("storage", "masterfile_schemas")


def sanitize_attribute_path(path: str) -> str:
    """Strips the characters Amazon's defined-name generator drops when it
    turns an attribute path like 'brand[marketplace_id=...]#1.value' into a
    workbook-safe name like 'brandmarketplace_id...1.value'."""
    return SANITIZE_RE.sub("", path or "")


def _find_template_sheet(wb):
    for name in TEMPLATE_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            ws = wb[name]
            a1 = ws.cell(row=1, column=1).value
            if a1 and "labelRow" in str(a1):
                return ws
    # Fall back to scanning every sheet for the same settings signature,
    # in case a different category's export names it something else.
    for name in wb.sheetnames:
        ws = wb[name]
        a1 = ws.cell(row=1, column=1).value
        if a1 and "labelRow" in str(a1) and "attributeRow" in str(a1):
            return ws
    raise ValueError(
        "Could not find the Template sheet (no sheet has the expected "
        "'labelRow=...&attributeRow=...' settings string in cell A1)."
    )


def _parse_settings(raw: str) -> dict:
    return dict(urllib.parse.parse_qsl(raw or "", keep_blank_values=True))


def _read_range_values(wb, ref_text: str) -> list[str]:
    """ref_text looks like "'Dropdown Lists'!$F$4:$F$183" -- reads every
    non-blank cell in that range, in order, de-duplicated."""
    if "!" not in ref_text:
        return []
    sheet_part, cell_part = ref_text.rsplit("!", 1)
    sheet_name = sheet_part.strip("'")
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(cell_part)
    seen = set()
    values = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            v = str(v).strip()
            if v and v not in seen:
                seen.add(v)
                values.append(v)
    return values


@dataclass
class FieldDef:
    attribute_path: str
    label: str
    group: str
    column_index: int
    required_level: str = "Optional"
    description: str = ""
    example: str = ""
    # product_type -> list of valid values; "_global" holds values for
    # attributes whose dropdown isn't product-type-specific.
    dropdown_values: dict = field(default_factory=dict)

    def values_for(self, product_type: str) -> list[str]:
        return self.dropdown_values.get(product_type) or self.dropdown_values.get("_global") or []

    def has_dropdown(self, product_type: str | None = None) -> bool:
        if product_type:
            return bool(self.values_for(product_type))
        return bool(self.dropdown_values)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "FieldDef":
        return cls(**d)


@dataclass
class TemplateSchema:
    source_filename: str
    sheet_name: str
    label_row: int
    attribute_row: int
    data_row: int
    group_row: int
    example_row: int
    product_types: list
    fields: list  # list[FieldDef]
    parsed_at: float = field(default_factory=time.time)

    def groups(self) -> list[str]:
        seen = []
        for f in self.fields:
            if f.group and f.group not in seen:
                seen.append(f.group)
        return seen

    def fields_by_group(self) -> dict:
        out: dict[str, list[FieldDef]] = {}
        for f in self.fields:
            out.setdefault(f.group or "Other", []).append(f)
        return out

    def to_dict(self) -> dict:
        d = asdict(self)
        return d

    @classmethod
    def from_dict(cls, d: dict) -> "TemplateSchema":
        d = dict(d)
        d["fields"] = [FieldDef.from_dict(f) for f in d["fields"]]
        return cls(**d)


def parse_masterfile(data: bytes, filename: str) -> TemplateSchema:
    """Parses one uploaded masterfile into a TemplateSchema. Raises
    ValueError with a plain-English message if it doesn't look like an
    Amazon flat-file template."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = _find_template_sheet(wb)

    settings = _parse_settings(ws.cell(row=1, column=1).value)
    label_row = int(settings.get("labelRow", DEFAULT_LABEL_ROW))
    attribute_row = int(settings.get("attributeRow", DEFAULT_ATTRIBUTE_ROW))
    data_row = int(settings.get("dataRow", DEFAULT_DATA_ROW))
    group_row = max(1, label_row - 1)
    example_row = max(1, data_row - 1)

    # Data Definitions: attribute_path -> (required_level, description, example)
    definitions: dict[str, tuple] = {}
    if DATA_DEFINITIONS_SHEET in wb.sheetnames:
        dd = wb[DATA_DEFINITIONS_SHEET]
        for r in range(1, dd.max_row + 1):
            attr = dd.cell(row=r, column=2).value
            if not attr:
                continue
            desc = dd.cell(row=r, column=4).value or ""
            example = dd.cell(row=r, column=5).value or ""
            required = dd.cell(row=r, column=6).value or "Optional"
            definitions[str(attr).strip()] = (str(required).strip(), str(desc).strip(), str(example).strip())

    # Walk the Template columns, carrying the group header forward across
    # the merged cells it spans.
    fields: list[FieldDef] = []
    current_group = ""
    product_type_col = None
    for c in range(1, ws.max_column + 1):
        g = ws.cell(row=group_row, column=c).value
        if g:
            current_group = str(g).strip()
        label = ws.cell(row=label_row, column=c).value
        attr = ws.cell(row=attribute_row, column=c).value
        example = ws.cell(row=example_row, column=c).value
        if not attr:
            continue
        attr = str(attr).strip()
        label = str(label).strip() if label else attr
        if label == "Product Type":
            product_type_col = c

        required_level, description, def_example = definitions.get(attr, ("Optional", "", ""))
        fields.append(FieldDef(
            attribute_path=attr,
            label=label,
            group=current_group,
            column_index=c,
            required_level=required_level,
            description=description,
            example=str(example).strip() if example else def_example,
        ))

    # Which product type(s) does this file actually contain data for?
    product_types: list[str] = []
    if product_type_col:
        seen = set()
        for r in range(data_row, ws.max_row + 1):
            v = ws.cell(row=r, column=product_type_col).value
            if v:
                v = str(v).strip()
                if v and v not in seen:
                    seen.add(v)
                    product_types.append(v)
        if not product_types:
            # Blank template with no data rows yet -- fall back to the
            # example row's product type, if there is one.
            v = ws.cell(row=example_row, column=product_type_col).value
            if v:
                product_types.append(str(v).strip())
    if not product_types:
        product_types = ["UNKNOWN"]

    # Resolve dropdown values for every field, for every product type found,
    # using Amazon's "{ProductType}{sanitized attribute path}" defined-name
    # convention (falling back to the bare sanitized path for attributes
    # that aren't product-type-specific).
    defined_names = set(wb.defined_names.keys())
    for f in fields:
        san = sanitize_attribute_path(f.attribute_path)
        if not san:
            continue
        for pt in product_types:
            key = f"{pt}{san}"
            if key in defined_names:
                ref = wb.defined_names[key].attr_text
                values = _read_range_values(wb, ref)
                if values:
                    f.dropdown_values[pt] = values
        if san in defined_names and "_global" not in f.dropdown_values:
            ref = wb.defined_names[san].attr_text
            values = _read_range_values(wb, ref)
            if values:
                f.dropdown_values["_global"] = values

    return TemplateSchema(
        source_filename=filename,
        sheet_name=ws.title,
        label_row=label_row,
        attribute_row=attribute_row,
        data_row=data_row,
        group_row=group_row,
        example_row=example_row,
        product_types=product_types,
        fields=fields,
    )


@dataclass
class MergeReport:
    product_type: str
    is_new: bool
    total_fields: int
    added_fields: list
    changed_fields: list
    storage_path: str


def _safe_filename(product_type: str) -> str:
    return re.sub(r"[^A-Za-z0-9_\-]", "_", product_type) + ".json"


def _schema_path(product_type: str) -> str:
    return os.path.join(STORAGE_DIR, _safe_filename(product_type))


def load_schema(product_type: str) -> TemplateSchema | None:
    path = _schema_path(product_type)
    if not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8") as fh:
        return TemplateSchema.from_dict(json.load(fh))


def list_stored_product_types() -> list[str]:
    if not os.path.isdir(STORAGE_DIR):
        return []
    out = []
    for name in sorted(os.listdir(STORAGE_DIR)):
        if name.endswith(".json"):
            try:
                with open(os.path.join(STORAGE_DIR, name), "r", encoding="utf-8") as fh:
                    d = json.load(fh)
                out.append(d.get("product_types", [name[:-5]])[0])
            except Exception:
                continue
    return out


def _merge_field(old: FieldDef, new: FieldDef) -> FieldDef:
    merged = FieldDef(
        attribute_path=old.attribute_path,
        label=new.label or old.label,
        group=new.group or old.group,
        column_index=new.column_index,
        required_level=new.required_level or old.required_level,
        description=new.description or old.description,
        example=new.example or old.example,
        dropdown_values=dict(old.dropdown_values),
    )
    for pt, values in new.dropdown_values.items():
        if values:
            merged.dropdown_values[pt] = values
    return merged


def save_schema_for_product_type(schema: TemplateSchema, product_type: str) -> MergeReport:
    """Persists (or merges into) the stored schema for a single product
    type found in `schema`. Fields already on record are updated in place
    (labels, required level, dropdown values); fields never seen before are
    added and listed in the report; nothing already on record is ever
    silently removed."""
    os.makedirs(STORAGE_DIR, exist_ok=True)
    existing = load_schema(product_type)

    added, changed = [], []

    if existing is None:
        new_schema = TemplateSchema(
            source_filename=schema.source_filename,
            sheet_name=schema.sheet_name,
            label_row=schema.label_row,
            attribute_row=schema.attribute_row,
            data_row=schema.data_row,
            group_row=schema.group_row,
            example_row=schema.example_row,
            product_types=[product_type],
            fields=[f for f in schema.fields],
        )
        added = [f.label for f in schema.fields]
        is_new = True
    else:
        by_path = {f.attribute_path: f for f in existing.fields}
        merged_fields = []
        for f in schema.fields:
            if f.attribute_path in by_path:
                old = by_path[f.attribute_path]
                merged = _merge_field(old, f)
                if merged.required_level != old.required_level or merged.dropdown_values != old.dropdown_values:
                    changed.append(merged.label)
                merged_fields.append(merged)
                del by_path[f.attribute_path]
            else:
                merged_fields.append(f)
                added.append(f.label)
        # Anything left in by_path existed before and wasn't in this
        # upload -- keep it, so a file that only shows a subset of columns
        # never erases history.
        merged_fields.extend(by_path.values())

        new_schema = TemplateSchema(
            source_filename=schema.source_filename,
            sheet_name=schema.sheet_name,
            label_row=schema.label_row,
            attribute_row=schema.attribute_row,
            data_row=schema.data_row,
            group_row=schema.group_row,
            example_row=schema.example_row,
            product_types=[product_type],
            fields=merged_fields,
        )
        is_new = False

    path = _schema_path(product_type)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(new_schema.to_dict(), fh, indent=2)

    return MergeReport(
        product_type=product_type,
        is_new=is_new,
        total_fields=len(new_schema.fields),
        added_fields=added,
        changed_fields=changed,
        storage_path=path,
    )


def save_schema(schema: TemplateSchema) -> list[MergeReport]:
    """A single uploaded file can carry more than one product type's data
    rows (same Template layout, different Product Type values) -- store
    each one separately."""
    return [save_schema_for_product_type(schema, pt) for pt in schema.product_types]